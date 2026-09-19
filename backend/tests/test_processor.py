from io import BytesIO

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from app.models import CaseStatus, EmailCategory, EmailRecord
from app.services.document_reader import document_to_text
from app.services.processor import CaseProcessor
from app.services.text_extractor import extract_shipping_fields

from tests.test_extractor import BL_TEXT, SI_TEXT


class FakeInbox:
    def __init__(self) -> None:
        self.email = EmailRecord(
            email_id="email_001",
            **{"from": "shipping@example.com"},
            subject="Please confirm SI and draft BL",
            body="Attached for checking.",
            attachments=["attachments/email_001_SI.txt", "attachments/email_001_BL.txt"],
        )

    async def list_emails(self) -> list[EmailRecord]:
        return [self.email]

    async def get_attachment(self, path: str) -> bytes:
        return SI_TEXT.encode() if "_SI." in path else BL_TEXT.encode()


class AttachmentInbox:
    def __init__(self, attachments: dict[str, bytes]) -> None:
        self.attachments = attachments
        self.email = EmailRecord(
            email_id="email_xlsx",
            **{"from": "shipping@example.com"},
            subject="Please confirm SI and draft BL",
            body="Attached for checking.",
            attachments=list(attachments),
        )

    async def list_emails(self) -> list[EmailRecord]:
        return [self.email]

    async def get_attachment(self, path: str) -> bytes:
        return self.attachments[path]


def make_xlsx(labels: tuple[str, ...]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    values: tuple[str | int, ...] = (
        "APRIL FAR EAST (M) SDN BHD",
        "MOORIM SP CO., LTD",
        "UAB NOVAKOPA",
        "PORT KLANG, MALAYSIA",
        "CALLAO, PERU",
        3,
        21577,
    )
    for row, (label, value) in enumerate(zip(labels, values, strict=True), start=4):
        worksheet.cell(row=row, column=1, value=label)
        worksheet.cell(row=row, column=2, value=value)
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


SI_XLSX = make_xlsx(
    (
        "Shipper/Exporter",
        "Consignee (Non-Negotiable)",
        "Notify Party",
        "Port of Loading",
        "Port of Discharge",
        "No. of Containers",
        "Gross Weight (KG)",
    )
)

BL_XLSX = make_xlsx(
    (
        "Shipper",
        "Consignee",
        "Notify",
        "POL",
        "POD",
        "Container Count",
        "Gross Wt (KGS)",
    )
)


DOCX_LABELS = (
    "Shipper (Principal or Seller) (发货人)",
    "Consignee (Non-Negotiable) (收货人)",
    "Notify Party/Intermediate Consignee (通知人)",
    "Port of Loading (装货港)",
    "Port of Discharge (POD) (卸货港)",
    "No. of Containers (箱数)",
    "Gross Weight毛重(KGS) (毛重 KGS)",
)

DOCX_VALUES: tuple[str | tuple[str, ...], ...] = (
    "APRIL FAR EAST (M) SDN BHD",
    "MOORIM SP CO., LTD",
    "UAB NOVAKOPA",
    "PORT KLANG, MALAYSIA",
    "CALLAO, PERU",
    "3",
    "21,577 KG",
)


def make_docx(values: tuple[str | tuple[str, ...], ...] = DOCX_VALUES) -> bytes:
    document = Document()
    table = document.add_table(rows=0, cols=2)
    for label, value in zip(DOCX_LABELS, values, strict=True):
        label_cell, value_cell = table.add_row().cells
        label_cell.text = label
        paragraphs = (value,) if isinstance(value, str) else value
        value_cell.text = paragraphs[0]
        for paragraph in paragraphs[1:]:
            value_cell.add_paragraph(paragraph)
    content = BytesIO()
    document.save(content)
    return content.getvalue()


BL_DOCX = make_docx()


PDF_ROWS: tuple[tuple[str, str], ...] = (
    ("Shipper", "APRIL FAR EAST (M) SDN BHD"),
    ("Consignee", "MOORIM SP CO., LTD"),
    ("Notify Party", "UAB NOVAKOPA"),
    ("Port of Loading", "PORT KLANG, MALAYSIA"),
    ("Port of Discharge", "CALLAO, PERU"),
    ("Container Count", "3 x 40'HC"),
    ("Gross Weight (KG)", "21,577 KG"),
)


def make_pdf(
    pages: tuple[tuple[tuple[str, str], ...], ...] = (PDF_ROWS,),
) -> bytes:
    document = pymupdf.open()
    for rows in pages:
        page = document.new_page()
        y = 72
        for label, value in rows:
            page.insert_text((56, y), label)
            page.insert_text((220, y), value)
            y += 24
    content = document.tobytes()
    document.close()
    return content


PDF_CONTENT = make_pdf()


def make_xlsx_inbox() -> AttachmentInbox:
    return AttachmentInbox(
        {
            "attachments/email_xlsx_SI.xlsx": SI_XLSX,
            "attachments/email_xlsx_BL.xlsx": BL_XLSX,
        }
    )


def make_xlsx_docx_inbox() -> AttachmentInbox:
    return AttachmentInbox(
        {
            "attachments/email_xlsx_SI.xlsx": SI_XLSX,
            "attachments/email_xlsx_BL.docx": BL_DOCX,
        }
    )


def make_txt_inbox(bl_text: str, si_text: str = SI_TEXT) -> AttachmentInbox:
    return AttachmentInbox(
        {
            "attachments/email_txt_SI.txt": si_text.encode(),
            "attachments/email_txt_BL.txt": bl_text.encode(),
        }
    )


async def test_processor_builds_a_complete_matching_case() -> None:
    processor = CaseProcessor(FakeInbox())
    cases = await processor.process_all()

    case = cases[0]
    assert case.category is EmailCategory.BL_COMPARISON
    assert case.status is CaseStatus.MATCH
    assert case.review_reason is None
    assert len(case.comparison) == 7


@pytest.mark.parametrize(
    "wrong_document",
    (
        "COMMERCIAL INVOICE\nSeller: ACME EXPORTS\nBuyer: ACME IMPORTS",
        "PACKING LIST\nShipper: ACME EXPORTS\nConsignee: ACME IMPORTS",
        "CERTIFICATE OF ORIGIN\nExporter: ACME EXPORTS",
        f"{BL_TEXT}\n*** THIS IS NOT AN SI OR BL ***",
        f"{BL_TEXT}\n*** PACKING LIST ONLY - NO PORT OR VESSEL DETAILS ***",
    ),
)
async def test_processor_marks_strong_wrong_document_indicators_for_review(
    wrong_document: str,
) -> None:
    inbox = make_txt_inbox(wrong_document)

    case = await CaseProcessor(inbox).process_email(inbox.email)

    assert case.status is CaseStatus.NEEDS_REVIEW
    assert case.review_reason.value == "wrong_doc_type"
    assert case.si_fields is None
    assert case.bl_fields is None
    assert case.comparison == []


async def test_processor_extracts_all_seven_fields_from_xlsx_si() -> None:
    inbox = make_xlsx_inbox()
    case = await CaseProcessor(inbox).process_email(inbox.email)

    assert case.status is CaseStatus.MATCH
    assert case.si_fields is not None
    assert all(value is not None for _, value in case.si_fields)
    assert case.si_fields.container_count.normalized_value == "3"
    assert case.si_fields.gross_weight_kg.normalized_value == "21577"


async def test_processor_extracts_all_seven_fields_from_xlsx_bl() -> None:
    inbox = make_xlsx_inbox()
    case = await CaseProcessor(inbox).process_email(inbox.email)

    assert case.status is CaseStatus.MATCH
    assert case.bl_fields is not None
    assert all(value is not None for _, value in case.bl_fields)
    assert case.bl_fields.container_count.normalized_value == "3"
    assert case.bl_fields.gross_weight_kg.normalized_value == "21577"


def test_document_to_text_reads_docx_table_and_strips_bilingual_suffixes() -> None:
    content = make_docx(
        (
            ("APRIL FAR EAST (M) SDN BHD", "80 RAFFLES PLACE"),
            *DOCX_VALUES[1:],
        )
    )

    text = document_to_text("attachments/email_xlsx_BL.docx", content)

    assert (
        "Shipper (Principal or Seller): "
        "APRIL FAR EAST (M) SDN BHD | 80 RAFFLES PLACE"
    ) in text
    assert "Notify Party/Intermediate Consignee: UAB NOVAKOPA" in text
    assert "Gross Weight毛重(KGS): 21,577 KG" in text
    assert "(发货人)" not in text
    assert "(通知人)" not in text


def test_docx_text_reuses_existing_extractor_for_all_seven_fields() -> None:
    fields = extract_shipping_fields(
        document_to_text("attachments/email_xlsx_BL.docx", BL_DOCX)
    )

    assert all(value is not None for _, value in fields)
    assert fields.container_count.normalized_value == "3"
    assert fields.gross_weight_kg.normalized_value == "21577"


async def test_processor_compares_xlsx_si_with_docx_bl() -> None:
    inbox = make_xlsx_docx_inbox()
    case = await CaseProcessor(inbox).process_email(inbox.email)

    assert case.status is CaseStatus.MATCH
    assert case.si_fields is not None
    assert case.bl_fields is not None
    assert len(case.comparison) == 7


def test_document_to_text_reads_positioned_pdf_as_canonical_rows() -> None:
    text = document_to_text("attachments/email_pdf_SI.pdf", PDF_CONTENT)

    assert "Shipper: APRIL FAR EAST (M) SDN BHD" in text
    assert "Port of Discharge: CALLAO, PERU" in text


def test_pdf_text_reuses_existing_extractor_for_all_seven_fields() -> None:
    fields = extract_shipping_fields(
        document_to_text("attachments/email_pdf_BL.pdf", PDF_CONTENT)
    )

    assert all(value is not None for _, value in fields)
    assert fields.container_count.normalized_value == "3"
    assert fields.gross_weight_kg.normalized_value == "21577"


def test_document_to_text_joins_multiple_pdf_pages() -> None:
    content = make_pdf((PDF_ROWS[:4], PDF_ROWS[4:]))

    fields = extract_shipping_fields(
        document_to_text("attachments/email_pdf_BL.pdf", content)
    )

    assert all(value is not None for _, value in fields)


async def test_processor_keeps_txt_extraction_behavior_unchanged() -> None:
    inbox = FakeInbox()
    case = await CaseProcessor(inbox).process_email(inbox.email)

    assert case.status is CaseStatus.MATCH
    assert case.si_fields is not None
    assert case.bl_fields is not None
    assert len(case.comparison) == 7


async def test_processor_marks_scanned_or_unreadable_pdf_for_review() -> None:
    blank_pdf = make_pdf(((),))
    inbox = AttachmentInbox(
        {
            "attachments/email_xlsx_SI.xlsx": SI_XLSX,
            "attachments/email_xlsx_BL.pdf": blank_pdf,
        }
    )

    case = await CaseProcessor(inbox).process_email(inbox.email)

    assert case.status is CaseStatus.NEEDS_REVIEW
    assert case.review_reason.value == "unreadable"
