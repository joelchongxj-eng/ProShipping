from io import BytesIO

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.main as main


SI_TEXT = b"""Shipper: ACME EXPORT LTD
Consignee: ACME IMPORT LTD
Notify Party: ACME NOTIFY LTD
Port of Loading: PORT KLANG
Port of Discharge: CALLAO
Container Count: 3
Gross Weight (KG): 21,577 KG
"""

BL_MISMATCH_AND_MISSING = b"""Shipper: DIFFERENT EXPORT LTD
Consignee: ACME IMPORT LTD
Port of Loading: PORT KLANG
Port of Discharge: CALLAO
Container Count: 3
Gross Weight (KG): 21,577 KG
"""

FIELD_ROWS = (
    ("Shipper", "ACME EXPORT LTD"),
    ("Consignee", "ACME IMPORT LTD"),
    ("Notify Party", "ACME NOTIFY LTD"),
    ("Port of Loading", "PORT KLANG"),
    ("Port of Discharge", "CALLAO"),
    ("Container Count", 3),
    ("Gross Weight (KG)", 21577),
)


def make_xlsx() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Shipping Data"
    for row, (label, value) in enumerate(FIELD_ROWS, start=1):
        worksheet.cell(row=row, column=1, value=label)
        worksheet.cell(row=row, column=2, value=value)
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def make_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=0, cols=2)
    for label, value in FIELD_ROWS:
        label_cell, value_cell = table.add_row().cells
        label_cell.text = str(label)
        value_cell.text = str(value)
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def make_pdf(rows: tuple[tuple[str, object], ...] = FIELD_ROWS) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    y = 72
    for label, value in rows:
        page.insert_text((56, y), str(label))
        page.insert_text((220, y), str(value))
        y += 24
    content = document.tobytes()
    document.close()
    return content


FORMAT_CASES = (
    pytest.param(
        "xlsx",
        make_xlsx(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "attachment",
        id="xlsx",
    ),
    pytest.param(
        "docx",
        make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "attachment",
        id="docx",
    ),
    pytest.param("pdf", make_pdf(), "application/pdf", "inline", id="pdf"),
)


@pytest.fixture(autouse=True)
def clear_upload_sessions() -> None:
    store = getattr(main, "upload_store", None)
    if store is not None:
        store.clear()
    yield
    if store is not None:
        store.clear()


def test_compare_upload_accepts_matching_txt_pair() -> None:
    response = TestClient(main.app).post(
        "/api/compare-upload",
        files=(
            ("si_file", ("shipping_instruction.txt", SI_TEXT, "text/plain")),
            ("bl_file", ("draft_bl.txt", SI_TEXT, "text/plain")),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "MATCH"
    assert payload["review_reason"] is None
    assert len(payload["comparison"]) == 7
    assert payload["si_file"]["filename"] == "shipping_instruction.txt"
    assert payload["bl_file"]["filename"] == "draft_bl.txt"
    comparison_id = payload["comparison_id"]
    assert payload["si_file"]["source_filename"].startswith(
        f"uploads/{comparison_id}/si/"
    )
    assert payload["si_file"]["attachment_url"] == (
        f"/api/upload-comparisons/{comparison_id}/attachments/si"
    )
    assert payload["si_fields"]["shipper"]["source"]["locator"]["kind"] == "txt"


def test_uploaded_comparison_and_originals_remain_retrievable() -> None:
    client = TestClient(main.app)
    response = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", ("shipping_instruction.txt", SI_TEXT, "text/plain")),
            ("bl_file", ("draft_bl.txt", SI_TEXT, "text/plain")),
        ),
    )
    payload = response.json()
    comparison_id = payload["comparison_id"]

    comparison_response = client.get(
        f"/api/upload-comparisons/{comparison_id}"
    )
    attachment_response = client.get(payload["si_file"]["attachment_url"])

    assert comparison_response.status_code == 200
    assert comparison_response.json() == payload
    assert attachment_response.status_code == 200
    assert attachment_response.content == SI_TEXT
    assert attachment_response.headers["content-type"].startswith("text/plain")
    assert attachment_response.headers["content-disposition"] == (
        'inline; filename="shipping_instruction.txt"'
    )


@pytest.mark.parametrize(
    ("extension", "content", "content_type", "disposition"),
    FORMAT_CASES,
)
def test_compare_upload_preserves_supported_format_source_metadata(
    extension: str,
    content: bytes,
    content_type: str,
    disposition: str,
) -> None:
    client = TestClient(main.app)
    response = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", (f"si.{extension}", content, content_type)),
            ("bl_file", (f"bl.{extension}", content, content_type)),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "MATCH"
    assert payload["si_fields"]["shipper"]["source"]["locator"]["kind"] == extension
    attachment = client.get(payload["si_file"]["attachment_url"])
    assert attachment.content == content
    assert attachment.headers["content-type"].startswith(content_type)
    assert attachment.headers["content-disposition"].startswith(disposition)


def test_compare_upload_preserves_mismatch_over_missing_precedence() -> None:
    response = TestClient(main.app).post(
        "/api/compare-upload",
        files=(
            ("si_file", ("si.txt", SI_TEXT, "text/plain")),
            ("bl_file", ("bl.txt", BL_MISMATCH_AND_MISSING, "text/plain")),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "MISMATCH"
    assert payload["review_reason"] == "missing_value"
    assert payload["comparison"][0]["status"] == "mismatch"
    assert payload["comparison"][2]["status"] == "missing"


@pytest.mark.parametrize(
    ("si_content", "si_filename", "expected_reason"),
    (
        (b"not a docx", "si.docx", "unreadable"),
        (b"COMMERCIAL INVOICE\nNOT AN SI OR BL", "si.txt", "wrong_doc_type"),
        (make_pdf(()), "si.pdf", "unreadable"),
    ),
)
def test_compare_upload_returns_needs_review_for_unprocessable_documents(
    si_content: bytes,
    si_filename: str,
    expected_reason: str,
) -> None:
    response = TestClient(main.app).post(
        "/api/compare-upload",
        files=(
            ("si_file", (si_filename, si_content, "application/octet-stream")),
            ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "NEEDS_REVIEW"
    assert payload["review_reason"] == expected_reason
    assert payload["comparison"] == []


@pytest.mark.parametrize(
    ("files", "detail"),
    (
        (
            (("bl_file", ("bl.txt", SI_TEXT, "text/plain")),),
            "Exactly one SI file is required.",
        ),
        (
            (("si_file", ("si.txt", SI_TEXT, "text/plain")),),
            "Exactly one BL file is required.",
        ),
        (
            (
                ("si_file", ("si.txt", SI_TEXT, "text/plain")),
                ("si_file", ("second.txt", SI_TEXT, "text/plain")),
                ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
            ),
            "Exactly one SI file is required.",
        ),
        (
            (
                ("si_file", ("si.txt", SI_TEXT, "text/plain")),
                ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
                ("bl_file", ("second.txt", SI_TEXT, "text/plain")),
            ),
            "Exactly one BL file is required.",
        ),
    ),
)
def test_compare_upload_requires_exactly_one_file_per_role(
    files: tuple[tuple[str, tuple[str, bytes, str]], ...],
    detail: str,
) -> None:
    response = TestClient(main.app).post("/api/compare-upload", files=files)

    assert response.status_code == 400
    assert response.json() == {"detail": detail}


@pytest.mark.parametrize(
    ("filename", "content", "status_code", "detail"),
    (
        ("si.csv", SI_TEXT, 400, "Unsupported SI file type."),
        ("si.txt", b"", 400, "SI file is empty."),
    ),
)
def test_compare_upload_rejects_invalid_file(
    filename: str,
    content: bytes,
    status_code: int,
    detail: str,
) -> None:
    response = TestClient(main.app).post(
        "/api/compare-upload",
        files=(
            ("si_file", (filename, content, "application/octet-stream")),
            ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
        ),
    )

    assert response.status_code == status_code
    assert response.json() == {"detail": detail}


def test_compare_upload_enforces_streaming_size_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main.upload_store, "max_upload_bytes", 4)

    response = TestClient(main.app).post(
        "/api/compare-upload",
        files=(
            ("si_file", ("si.txt", b"12345", "text/plain")),
            ("bl_file", ("bl.txt", b"1234", "text/plain")),
        ),
    )

    assert response.status_code == 413
    assert response.json() == {
        "detail": "SI file exceeds the upload size limit."
    }


def test_compare_upload_sanitizes_client_filename() -> None:
    client = TestClient(main.app)
    response = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", ("../../shipping_instruction.txt", SI_TEXT, "text/plain")),
            ("bl_file", ("..\\draft_bl.txt", SI_TEXT, "text/plain")),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["si_file"]["filename"] == "shipping_instruction.txt"
    assert payload["bl_file"]["filename"] == "draft_bl.txt"
    assert ".." not in payload["si_file"]["source_filename"]
    assert client.get(payload["si_file"]["attachment_url"]).content == SI_TEXT


def test_compare_upload_sanitizes_header_control_characters() -> None:
    client = TestClient(main.app)
    response = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", ('evil"name.txt', SI_TEXT, "text/plain")),
            ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["si_file"]["filename"] == "evil_name.txt"
    attachment = client.get(payload["si_file"]["attachment_url"])
    assert attachment.headers["content-disposition"] == (
        'inline; filename="evil_name.txt"'
    )


def test_same_client_filename_does_not_overwrite_the_other_role() -> None:
    client = TestClient(main.app)
    response = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", ("document.txt", SI_TEXT, "text/plain")),
            (
                "bl_file",
                ("document.txt", BL_MISMATCH_AND_MISSING, "text/plain"),
            ),
        ),
    )

    assert response.status_code == 200
    payload = response.json()
    assert client.get(payload["si_file"]["attachment_url"]).content == SI_TEXT
    assert client.get(payload["bl_file"]["attachment_url"]).content == (
        BL_MISMATCH_AND_MISSING
    )


def test_upload_comparison_expires_after_configured_ttl(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main.upload_store, "ttl_seconds", 0)
    client = TestClient(main.app)
    response = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", ("si.txt", SI_TEXT, "text/plain")),
            ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
        ),
    )

    assert response.status_code == 200
    comparison_id = response.json()["comparison_id"]
    assert client.get(f"/api/upload-comparisons/{comparison_id}").status_code == 404


def test_upload_comparison_returns_404_for_unknown_id_or_role() -> None:
    client = TestClient(main.app)

    created = client.post(
        "/api/compare-upload",
        files=(
            ("si_file", ("si.txt", SI_TEXT, "text/plain")),
            ("bl_file", ("bl.txt", SI_TEXT, "text/plain")),
        ),
    ).json()

    unknown = client.get("/api/upload-comparisons/unknown")
    unknown_attachment = client.get(
        "/api/upload-comparisons/unknown/attachments/si"
    )
    invalid_role = client.get(
        f"/api/upload-comparisons/{created['comparison_id']}/attachments/other"
    )

    assert unknown.status_code == 404
    assert unknown.json() == {"detail": "Upload comparison not found."}
    assert unknown_attachment.status_code == 404
    assert unknown_attachment.json() == {
        "detail": "Uploaded attachment not found."
    }
    assert invalid_role.status_code == 404
    assert invalid_role.json() == {"detail": "Uploaded attachment not found."}
