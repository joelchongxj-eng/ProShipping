from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class DocumentReadError(Exception):
    """The attachment format is unsupported or its content cannot be read."""


def _cell_to_text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _xlsx_to_text(content: bytes) -> str:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise DocumentReadError("Unable to read XLSX attachment.") from exc

    try:
        worksheet = workbook.active
        lines: list[str] = []
        for label, value in worksheet.iter_rows(min_col=1, max_col=2, values_only=True):
            if label is None or value is None:
                continue
            label_text = _cell_to_text(label)
            value_text = _cell_to_text(value)
            if label_text and value_text:
                lines.append(f"{label_text}: {value_text}")
        return "\n".join(lines)
    finally:
        workbook.close()


def document_to_text(filename: str, content: bytes) -> str:
    suffix = PurePosixPath(filename).suffix.casefold()
    if suffix == ".txt":
        return content.decode("utf-8-sig")
    if suffix == ".xlsx":
        return _xlsx_to_text(content)
    raise DocumentReadError(f"Unsupported document format: {suffix or '<none>'}")
